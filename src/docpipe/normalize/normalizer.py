"""Converts every supported input format into ONE standard processing format:
a NormalizedDocument with plain text (+ a needs_ocr flag for scans/images).
Structured files (CSV/JSON/XLSX) become 'Label: value' lines so one extractor
can handle all of them."""
import csv
import io
import json
import logging
from pathlib import Path

from ..models import NormalizedDocument, PermanentError, RawDocument

log = logging.getLogger(__name__)
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def _pdf(content: bytes) -> NormalizedDocument:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((p.extract_text() or "") for p in reader.pages).strip()
        pages = len(reader.pages)
    except Exception as exc:
        raise PermanentError(f"Corrupt or unreadable PDF: {exc}") from exc
    # No text layer => scanned PDF => needs OCR (Azure Document Intelligence)
    return NormalizedDocument("pdf", text, pages, needs_ocr=not text)


def _docx(content: bytes) -> NormalizedDocument:
    import docx
    try:
        d = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise PermanentError(f"Corrupt Word file: {exc}") from exc
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) == 2:
                lines.append(f"{cells[0]}: {cells[1]}")
    return NormalizedDocument("docx", "\n".join(lines))


def _xlsx(content: bytes) -> NormalizedDocument:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise PermanentError(f"Corrupt Excel file: {exc}") from exc
    ws = wb.active
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    return NormalizedDocument("xlsx", _rows_to_text(rows))


def _csv(content: bytes) -> NormalizedDocument:
    try:
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    except Exception as exc:
        raise PermanentError(f"Unreadable CSV: {exc}") from exc
    return NormalizedDocument("csv", _rows_to_text(rows))


def _rows_to_text(rows: list) -> str:
    """Header row + first data row -> 'header: value'. Two-column sheets are
    treated as key/value pairs."""
    if not rows:
        return ""
    if len(rows[0]) == 2 and len(rows) > 2:
        return "\n".join(f"{r[0]}: {r[1]}" for r in rows if r[0] is not None)
    if len(rows) < 2:
        raise PermanentError("Spreadsheet has a header but no data row")
    return "\n".join(f"{h}: {v}" for h, v in zip(rows[0], rows[1]) if h is not None)


def _json(content: bytes) -> NormalizedDocument:
    try:
        data = json.loads(content.decode("utf-8"))
    except Exception as exc:
        raise PermanentError(f"Invalid JSON: {exc}") from exc
    if isinstance(data, list) and data:
        data = data[0]
    if not isinstance(data, dict):
        raise PermanentError("JSON root must be an object or a list of objects")
    return NormalizedDocument("json", "\n".join(f"{k}: {v}" for k, v in data.items()
                                                  if not isinstance(v, (dict, list))))


def normalize(doc: RawDocument) -> NormalizedDocument:
    ext = Path(doc.filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(doc.content)
    if ext == ".docx":
        return _docx(doc.content)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(doc.content)
    if ext == ".csv":
        return _csv(doc.content)
    if ext == ".json":
        return _json(doc.content)
    if ext == ".txt":
        return NormalizedDocument("txt", doc.content.decode("utf-8", errors="replace"))
    if ext in IMAGE_EXT:
        return NormalizedDocument("image", "", needs_ocr=True)
    raise PermanentError(f"Unsupported file type: {ext or 'none'}")
