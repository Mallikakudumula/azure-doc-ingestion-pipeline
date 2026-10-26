"""Creates 100% FAKE sample documents in data/inbox/<channel>/ to demo every path:
completed, needs-review, duplicate, idempotent skip, dead-letter."""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
INBOX = ROOT / "data" / "inbox"


def out(channel: str, name: str) -> Path:
    p = INBOX / channel / name
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def make_pdf(path: Path, lines: list[str]) -> None:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), invariant=1)    # invariant => identical bytes each run
    y = 800
    for line in lines:
        c.drawString(60, y, line)
        y -= 22
    c.save()


def main() -> None:
    inv = ["INVOICE", "Invoice Number: INV-1001", "Vendor: Acme Corp", "Date: 2025-03-14",
           "Total: 12,450.00", "Currency: USD"]
    # 1. clean PDF from email -> COMPLETED
    make_pdf(out("email", "invoice_acme.pdf"), inv)
    # 2. same file arrives again via shared drive -> idempotent SKIP (same hash)
    make_pdf(out("shared_drive", "invoice_acme_copy.pdf"), inv)
    # 3. same business invoice, different bytes -> DUPLICATE (business-key check)
    out("portal", "invoice_acme_resubmitted.txt").write_text(
        "Invoice Number: INV-1001\nVendor: Acme Corp\nInvoice Date: 14/03/2025\nTotal: 12450.00\nCurrency: USD\n")
    # 4. bad OCR-like amount -> low confidence -> NEEDS_REVIEW
    make_pdf(out("email", "blurry_invoice.pdf"),
             ["INVOICE", "Invoice Number: INV-1002", "Vendor: Initech", "Date: 2025-04-02",
              "Total: 8,9O0.00", "Currency: USD"])
    # 5. Word loan agreement with an account number -> Highly Confidential
    import docx
    d = docx.Document()
    d.add_heading("Loan Agreement", 1)
    t = d.add_table(rows=0, cols=2)
    for k, v in [("Loan ID", "LN-2002"), ("Borrower", "Globex Ltd"), ("Agreement Date", "2025-05-20"),
                 ("Loan Amount", "250000"), ("Currency", "EUR"), ("Account Number", "123456789012")]:
        r = t.add_row().cells
        r[0].text, r[1].text = k, v
    d.save(out("sharepoint", "loan_agreement.docx"))
    # 6. Excel trade sheet
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Reference", "Counterparty", "Date", "Amount", "Currency"])
    ws.append(["TR-3003", "Umbrella Bank", "2025-06-01", 98000.5, "GBP"])
    wb.save(out("sharepoint", "trade_sheet.xlsx"))
    # 7. CSV payment
    out("shared_drive", "payment.csv").write_text(
        "Invoice Number,Vendor,Date,Total,Currency\nINV-4004,Hooli,2025-07-11,1500.75,USD\n")
    # 8. JSON from the portal with a missing field -> NEEDS_REVIEW (completeness)
    out("portal", "deal_missing_date.json").write_text(json.dumps(
        {"document_number": "DL-5005", "counterparty": "Soylent plc", "amount": 42000, "currency": "USD"}))
    # 9. corrupt PDF -> DEAD LETTER
    out("shared_drive", "corrupt.pdf").write_bytes(b"this is not really a pdf")
    # 10. scanned image -> needs real OCR (works in azure mode; dead-letters in local mock mode)
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(img).text((20, 80), "Invoice Number: INV-6006  Vendor: Pied Piper  Total: 500.00", fill="black")
    img.save(out("portal", "scanned_invoice.png"))
    print(f"Sample documents created under {INBOX}")


if __name__ == "__main__":
    sys.exit(main())
